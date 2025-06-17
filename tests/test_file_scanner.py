# Add src directory to sys.path so we can import FileScanner
import sys
import os
import shutil
sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), '..', 'src')))
from file_scanner import FileScanner  # Import the FileScanner class from the src directory

import unittest  # Import unittest framework for testing
from PIL import Image  # For creating a sample image file
import docx  # For creating a sample docx file
from fpdf import FPDF  # For creating a sample PDF file

class TestFileScanner(unittest.TestCase):  # Define a test case class for FileScanner
    def setUp(self):  # Setup runs before each test
        self.scanner = FileScanner()  # Create an instance of FileScanner
        self.test_dir = os.path.join(os.path.dirname(__file__), 'test_files')  # Path for test files directory
        os.makedirs(self.test_dir, exist_ok=True)  # Create the test directory if it doesn't exist

        # Create a sample text file
        self.txt_expected = 'This is a credit card test text file.'
        self.txt_path = os.path.join(self.test_dir, 'sample.txt')  # Path for sample text file
        with open(self.txt_path, 'w', encoding='utf-8') as f:  # Open the text file for writing
            f.write(self.txt_expected)  # Write sample content to the text file

        # Create a sample docx file
        self.docx_expected = 'This is a public test docx file.'
        self.docx_path = os.path.join(self.test_dir, 'sample.docx')  # Path for sample docx file
        doc = docx.Document()  # Create a new docx document
        doc.add_paragraph(self.docx_expected)  # Add a paragraph to the docx file
        doc.save(self.docx_path)  # Save the docx file

        # Create a sample PDF file (text-based)
        self.pdf_expected = 'This is a do not distribute test PDF file.'
        self.pdf_path = os.path.join(self.test_dir, 'sample.pdf')  # Path for sample PDF file
        pdf = FPDF()  # Create a new PDF document
        pdf.add_page()  # Add a page to the PDF
        pdf.set_font('Arial', size=12)  # Set the font for the PDF
        pdf.cell(200, 10, self.pdf_expected, ln=True)  # Add a cell with text to the PDF
        pdf.output(self.pdf_path)  # Save the PDF file
        
        # Create a sample image file (with text)
        self.img_expected = 'CONFIDENTIAL contents included in this image.'
        self.img_path = os.path.join(self.test_dir, 'sample.png')  # Path for sample image file
        img = Image.new('RGB', (200, 60), color=(255, 255, 255))  # Create a new blank image
        try:
            from PIL import ImageDraw, ImageFont  # Import drawing tools
            d = ImageDraw.Draw(img)  # Create a drawing context
            d.text((10, 10), self.img_expected, fill=(0, 0, 0))  # Draw text on the image
        except Exception:
            pass  # If drawing tools are not available, leave the image blank
        img.save(self.img_path)  # Save the image file

        # Create a sample PDF file with an image containing text
        from io import BytesIO
        from PIL import ImageDraw
        self.pdf_img_expected = 'CONFIDENTIAL scanned pdf. Do not share.'
        self.pdf_img_path = os.path.join(self.test_dir, 'sample_img.pdf')  # Path for sample PDF with image
        # Create an image with text
        img2 = Image.new('RGB', (200, 60), color=(255, 255, 255))
        d2 = ImageDraw.Draw(img2)
        d2.text((10, 10), self.pdf_img_expected, fill=(0, 0, 0))
        img_bytes = BytesIO()
        img2.save(img_bytes, format='PNG')
        img_bytes.seek(0)
        # Create a PDF and embed the image
        pdf2 = FPDF()
        pdf2.add_page()
        pdf2.set_font('Arial', size=12)
        # Save the image to a temporary file to insert into PDF
        temp_img_path = os.path.join(self.test_dir, 'temp_img.png')
        img2.save(temp_img_path)
        pdf2.image(temp_img_path, x=10, y=20, w=60)
        pdf2.output(self.pdf_img_path)
        os.remove(temp_img_path)

    # def tearDown(self):  # Teardown runs after each test
    #     # Prompt the user to decide whether to clean up the test directory
    #     response = input(f"Do you want to delete the test directory '{self.test_dir}' and its contents? (y/n): ").strip().lower()
    #     if response == 'y':
    #         if os.path.exists(self.test_dir):  # Check if the test directory exists
    #             shutil.rmtree(self.test_dir)  # Remove the test directory and all its contents
    #     else:
    #         print(f"Test directory '{self.test_dir}' was not deleted.")

    #def test_scan_directory(self):  # Test scanning the directory for all file types
    #    files = self.scanner.scan_directory(self.test_dir)  # Scan the test directory
    #    self.assertTrue(any('sample.txt' in f['path'] for f in files))  # Assert that the text file is found
    #    self.assertTrue(any('sample.docx' in f['path'] for f in files))  # Assert that the docx file is found
    #    self.assertTrue(any('sample.pdf' in f['path'] for f in files))  # Assert that the PDF file is found
    #    self.assertTrue(any('sample.png' in f['path'] for f in files))  # Assert that the image file is found

    def test_read_txt(self):  # Test reading a text file
        content = self.scanner.read_file(self.txt_path)  # Read the text file
        #print("!!!!!!")
        #print(self.txt_expected)
        expected = self.txt_expected
        self.assertEqual(content, expected)  # Assert the content matches what was written

    def test_read_docx(self):  # Test reading a docx file
        content = self.scanner.read_file(self.docx_path)  # Read the docx file
        expected = self.docx_expected
        self.assertIn(expected, content)  # Assert the content includes the expected text

    def test_read_pdf(self):  # Test reading a PDF file
        content = self.scanner.read_file(self.pdf_path)  # Read the PDF file
        expected = self.pdf_expected
        self.assertIn(expected, content)  # Assert the content includes the expected text

    def test_read_image(self):  # Test reading an image file with OCR
        content = self.scanner.read_file(self.img_path)  # Read the image file
        self.assertIsInstance(content, str)  # Assert the result is a string
        expected = self.img_expected
        from difflib import SequenceMatcher
        percent_match = SequenceMatcher(None, expected.lower(), content.lower()).ratio() * 100
        print(f"Image OCR percent match: {percent_match:.2f}% (Expected: '{expected}', Got: '{content.strip()}')")
        self.assertGreater(percent_match, 70, f"OCR percent match too low: {percent_match:.2f}%")  # Accept if >70%

    def test_read_pdf_with_image(self):  # Test reading a PDF file with an embedded image
        content = self.scanner.read_file(self.pdf_img_path)  # Read the PDF file with image
        self.assertIsInstance(content, str)  # Assert the result is a string
        expected = self.pdf_img_expected
        from difflib import SequenceMatcher
        percent_match = SequenceMatcher(None, expected.lower(), content.lower()).ratio() * 100
        print(f"PDF OCR percent match: {percent_match:.2f}% (Expected: '{expected}', Got: '{content.strip()}')")
        self.assertGreater(percent_match, 70, f"PDF OCR percent match too low: {percent_match:.2f}%")  # Accept if >70%

if __name__ == '__main__':  # Run the tests if this file is executed directly
    unittest.main()  # Start the unittest test runner